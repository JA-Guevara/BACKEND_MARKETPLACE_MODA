"""Exportacion de reportes analiticos (CSV tabular y XLSX de trabajo con hoja
de criterios). Filas y totales deben coincidir con el dashboard: el exportador
solo serializa el conjunto que ya fue filtrado y validado por ReportsService."""
import csv
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from src.ventas_pagos.application.reports_service import BUSINESS_TZ


def _safe(value):
    """Evita que un texto sea interpretado como formula por la planilla
    (=, +, -, @, tab) manteniendo el dato legible."""
    if not isinstance(value, str):
        return value
    if value and value[0] in {"=", "+", "-", "@", "\t", "\r"}:
        return "'" + value
    return value


def to_csv(title: str, headers: list[str], rows: list[dict], metadata: dict) -> bytes:
    """CSV con comillas y escapado correcto (RFC 4180): un nombre como
    "Camisa, lino" no rompe columnas porque el modulo csv cita/escapa."""
    buf = io.StringIO()
    order = list(headers)
    label = lambda key: key.replace("_", " ").capitalize()
    buf.write("# " + title + "\n")
    for key in ("periodo", "filtros", "zona", "moneda", "generado", "nota"):
        buf.write(f"# {label(key)}: {metadata.get(key, '-')}\n")
    writer = csv.writer(buf, lineterminator="\n")
    writer.writerow(headers)
    for row in rows:
        writer.writerow([_safe(row.get(key, "")) for key in order])
    return buf.getvalue().encode("utf-8-sig")


def to_xlsx(title: str, headers: list[str], rows: list[dict], metadata: dict) -> io.BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Reporte"
    sheet.freeze_panes = "A2"
    header_fill = PatternFill("solid", fgColor="74394E")
    header_font = Font(color="FFFFFF", bold=True)

    sheet.append([label(key) for key in headers])
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    for row in rows:
        sheet.append([safe_value(row.get(key, "")) for key in headers])

    # Anchos por columna. Con filas vacias se usa como referencia la cabecera;
    # asi el export no revienta con conjuntos filtrados a cero filas.
    widths_by_key: dict[str, int] = {}
    for row in rows:
        for key in headers:
            widths_by_key[key] = max(widths_by_key.get(key, 0), len(str(row.get(key, ""))))
    widths = [max(widths_by_key.get(key, 0), len(label(key))) + 2 for key in headers]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = min(width, 42)

    criteria = workbook.create_sheet("Criterios")
    criteria["A1"] = "Reporte"
    criteria["B1"] = title
    row_index = 2
    for key in ("periodo", "filtros", "zona", "moneda", "generado", "nota"):
        criteria.cell(row=row_index, column=1, value=label(key))
        criteria.cell(row=row_index, column=2, value=metadata.get(key, "-"))
        row_index += 1
    criteria.column_dimensions["A"].width = 18
    criteria.column_dimensions["B"].width = 70

    output = io.BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def label(key: str) -> str:
    return key.replace("_", " ").capitalize()


def safe_value(value):
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, datetime):
        return value.isoformat()
    return _safe(value)


def build_metadata(title: str, periodo: str, filtros: str, currency: str, truncated: bool = False) -> dict:
    nota = title + " exportado desde el dashboard de FashionStore."
    if truncated:
        nota += " Se aplico el limite maximo de filas configurado; el total del dashboard puede diferir."
    return {
        "periodo": periodo,
        "filtros": filtros or "sin filtros adicionales",
        "zona": "America/La_Paz",
        "moneda": currency,
        "generado": datetime.now(BUSINESS_TZ).strftime("%Y-%m-%d %H:%M %Z"),
        "nota": nota,
    }