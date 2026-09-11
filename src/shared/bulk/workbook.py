"""Bounded XLSX parsing; formula cells and external workbook links are forbidden."""
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from src.shared.exceptions.domain_exception import ValidationError

MAX_BYTES = 5 * 1024 * 1024
MAX_EXPANDED_BYTES = 25 * 1024 * 1024
MAX_ROWS = 1000
MAX_EXPORT_ROWS = 10000


def read_rows(content: bytes, columns: tuple[str, ...]):
    if not content or len(content) > MAX_BYTES:
        raise ValidationError('El archivo debe pesar como máximo 5 MB.')
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 150 or sum(x.file_size for x in entries) > MAX_EXPANDED_BYTES:
                raise ValidationError('El archivo descomprimido excede el límite permitido.')
            if any('vbaProject' in x.filename or 'externalLinks/' in x.filename for x in entries):
                raise ValidationError('No se admiten macros ni vínculos externos.')
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            if 'Datos' not in workbook.sheetnames:
                raise ValidationError('Falta la hoja Datos. Descargá la plantilla de este recurso.')
            sheet = workbook['Datos']
            if (sheet.max_row or 0) > MAX_ROWS + 1 or (sheet.max_column or 0) > len(columns):
                raise ValidationError('Máximo 1000 filas y las columnas originales de la plantilla.')
            rows = sheet.iter_rows(max_row=MAX_ROWS + 2, max_col=len(columns) + 1)
            header = next(rows, ())
            names = tuple(str(c.value).strip() if c.value is not None else '' for c in header)
            while names and not names[-1]:
                names = names[:-1]
            if names != columns:
                raise ValidationError('Las columnas deben coincidir con la plantilla: ' + ', '.join(columns))
            result = []
            for number, cells in enumerate(rows, 2):
                if not any(c.value is not None for c in cells):
                    continue
                if number > MAX_ROWS + 1:
                    raise ValidationError('Máximo 1000 filas por importación.')
                if any(c.data_type == 'f' for c in cells):
                    raise ValidationError(f'Fila {number}: no se admiten fórmulas; pegá solamente valores.')
                if any(isinstance(c.value, str) and len(c.value) > 10000 for c in cells):
                    raise ValidationError(f'Fila {number}: una celda supera los 10000 caracteres.')
                values = {name: cell.value.strip() if isinstance(cell.value, str) else cell.value
                          for name, cell in zip(columns, cells)}
                result.append((number, values))
            if not result:
                raise ValidationError('La hoja Datos no contiene registros.')
            return result
        finally:
            workbook.close()
    except ValidationError:
        raise
    except (BadZipFile, KeyError, ValueError, OSError, TypeError, StopIteration) as exc:
        raise ValidationError('Archivo Excel inválido. Utilizá la plantilla .xlsx.') from exc


def _cell_value(value):
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool, date)):
        return value
    if isinstance(value, Decimal):
        return float(value)
    return str(value)


def build_workbook(columns, rows, lookups, instructions):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Datos'
    sheet.append(list(columns))
    for row in rows:
        sheet.append([_cell_value(row.get(name)) for name in columns])
    guide = workbook.create_sheet('Instrucciones')
    for instruction in instructions:
        guide.append([instruction])
    for name, values in lookups.items():
        lookup = workbook.create_sheet(name[:31])
        lookup.append(['Valor permitido', 'Descripción'])
        for key, label in values:
            lookup.append([key, label])
    for page in workbook:
        page.freeze_panes = 'A2'
        page.auto_filter.ref = page.dimensions
        for cell in page[1]:
            cell.fill = PatternFill('solid', fgColor='183D3D')
            cell.font = Font(color='FFFFFF', bold=True)
        for row in page:
            for cell in row:
                # Explicit XLSX string cells keep =,+,-,@ text inert and preserve leading zeroes.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                    cell.number_format = '@'
        for column in page.columns:
            page.column_dimensions[column[0].column_letter].width = min(65, max(18, len(str(column[0].value or '')) + 4))
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
