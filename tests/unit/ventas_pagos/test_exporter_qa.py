"""Regresiones del exportador: CSV con comas/comillas/saltos/formulas, XLSX con
cero filas (regresion del UnboundLocalError) y nota de truncamiento."""
import io

from openpyxl import load_workbook

from src.ventas_pagos.application.exporter import build_metadata, to_csv, to_xlsx


def test_csv_escapa_comas_comillas_y_saltos():
    headers = ["nombre", "nota"]
    rows = [
        {"nombre": "Camisa, lino", "nota": "hola"},
        {"nombre": 'Cita "doble"', "nota": "multi\nlinea"},
        {"nombre": "=INDIRECT(A1)", "nota": ""},
    ]
    metadata = {"periodo": "todo", "filtros": "", "zona": "America/La_Paz", "moneda": "BOB", "generado": "ahora"}
    text = to_csv("Ventas", headers, rows, metadata).decode("utf-8-sig")
    assert '"Camisa, lino"' in text          # la coma NO rompe la columna
    assert '"Cita ""doble"""' in text        # comillas escapadas
    assert 'multi\nlinea' in text            # el salto queda dentro de la celda
    assert "'=INDIRECT(A1)" in text          # formula neutralizada


def test_csv_registra_metadata():
    metadata = build_metadata("Ventas", "2025-01-01 a 2025-01-31", "sin filtros", "BOB")
    text = to_csv("Ventas", ["fecha", "total"], [], metadata).decode("utf-8-sig")
    assert "# Periodo: 2025-01-01 a 2025-01-31" in text
    assert "# Moneda: BOB" in text


def test_xlsx_soporta_conjunto_vacio():
    metadata = build_metadata("Ventas", "todo", "", "BOB")
    output = to_xlsx("Ventas", ["fecha", "total"], [], metadata)
    workbook = load_workbook(io.BytesIO(output.getvalue()))
    sheet = workbook["Reporte"]
    assert sheet["A1"].value == "Fecha"
    assert sheet["B1"].value == "Total"
    assert sheet.max_row == 1
    criteria = workbook["Criterios"]
    assert criteria["B1"].value == "Ventas"


def test_xlsx_celda_texto_numero_y_formula():
    headers = ["nombre", "total", "fecha"]
    rows = [{"nombre": "=SUM(A1:A9)", "total": 50.0, "fecha": "2025-01-02T12:00:00"}]
    output = to_xlsx("Ventas", headers, rows, build_metadata("Ventas", "todo", "", "BOB"))
    sheet = load_workbook(io.BytesIO(output.getvalue()))["Reporte"]
    assert sheet["A2"].value == "'=SUM(A1:A9)"   # formula neutralizada
    assert sheet["B2"].value == 50.0             # numero real


def test_metadata_advierte_sobre_truncamiento():
    metadata = build_metadata("Ventas", "todo", "", "BOB", truncated=True)
    assert "limite maximo de filas" in metadata["nota"]
    metadata = build_metadata("Ventas", "todo", "", "BOB")
    assert "limite maximo de filas" not in metadata["nota"]