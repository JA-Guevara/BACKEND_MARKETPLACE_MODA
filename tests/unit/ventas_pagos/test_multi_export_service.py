"""Contrato de la exportacion multiple: coherencia dashboard=export, limites de
filas con aviso de truncamiento, csv unico vs zip y resumen para bitacora."""
import io
import zipfile
import uuid
from datetime import datetime, timedelta, timezone
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import Session

from src.auth.infrastructure.persistence.models.user import UserModel
from src.infrastructure.database.base import Base
from src.inventario_sucursales.application.services.organization_service import OrganizationService
from src.inventario_sucursales.web.schemas.organization import BranchCreate, CityCreate
from src.usuarios_catalogo.application.services.catalog_service import CatalogService
from src.usuarios_catalogo.web.schemas.catalog import CategoryCreate, ColorCreate, ProductCreate, SizeCreate, VariantCreate
from src.ventas_pagos.application.multi_export_service import MultiExportService
from src.ventas_pagos.application.reports_service import ReportsService
from src.ventas_pagos.infrastructure.models import OrderModel, StockModel

UTC = timezone.utc


def make_world() -> tuple[Session, dict]:
    engine = create_engine("sqlite+pysqlite:///:memory:")
    Base.metadata.create_all(engine)
    db = Session(engine, expire_on_commit=False)
    actor = UserModel(email="qa@fashionstore.test", password_hash="x", first_name="QA", last_name="Tienda",
                      is_active=True, is_verified=True)
    db.add(actor)
    db.flush()
    org = OrganizationService(db)
    city = org.create_city(CityCreate(name="Santa Cruz", department="Santa Cruz"), actor)
    branch = org.create_branch(BranchCreate(code="SCZ-01", name="Sucursal Central", city_id=city.id, address="Av. 1"), actor)
    catalog = CatalogService(db)
    cat_a = catalog.create_category(CategoryCreate(name="Camisas"), actor)
    size = catalog.create_size(SizeCreate(code="M", name="Mediana", sort_order=20), actor)
    color = catalog.create_color(ColorCreate(name="Rojo", hex_code="#FF0000"), actor)
    product_a = catalog.create_product(ProductCreate(
        name="Camisa, lino", description="descripcion de QA", brand="FashionStore", gender="mujer",
        base_price=Decimal("50"), category_id=cat_a.id,
        variants=[VariantCreate(size_id=size.id, color_id=color.id, sku="CAM-001")]), actor)
    va = product_a.variants[0]
    db.add(StockModel(variant_id=va.id, branch_id=branch.id, quantity=20))
    db.commit()
    return db, {"actor": actor, "branch": branch, "cat_a": cat_a, "va": va, "product_a": product_a}


def line(variant, product, name, unit_price, quantity):
    return {"variant_id": str(variant.id), "product_id": str(product.id), "name": name,
            "quantity": quantity, "unit_price": str(unit_price), "line_total": str(unit_price * quantity)}


def add_order(db, branch, user, items, created_at, paid_at=None):
    order = OrderModel(number=uuid.uuid4().hex[:24].upper()[:24], user_id=user.id, customer_email="c@mail.test",
                       branch_id=branch.id, status="paid", payment_status="paid", payment_method="stripe",
                       total=sum(Decimal(i["line_total"]) for i in items), currency="bob", address={},
                       items=items, tracking=[], paid_at=paid_at, created_at=created_at, updated_at=created_at)
    db.add(order)
    db.commit()
    return order


def test_export_multiple_coincide_con_dashboard():
    db, world = make_world()
    u, b = world["actor"], world["branch"]
    va, pa = world["va"], world["product_a"]
    now = datetime.now(UTC)
    add_order(db, b, u, [line(va, pa, "Camisa, lino", 50, 2)], now - timedelta(days=1), paid_at=now - timedelta(days=1))

    data = ReportsService(db).dashboard()
    service = MultiExportService(db)
    payload, media_type, filename, summary = service.export(["ventas", "pedidos", "existencias"], "xlsx")
    assert media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert filename.endswith(".xlsx")
    workbook = load_workbook(io.BytesIO(payload.getvalue()))
    assert {"Criterios", "Ventas", "Pedidos", "Existencias"} <= set(workbook.sheetnames)
    ventas_total = sum(r[2] for r in workbook["Ventas"].iter_rows(min_row=2, values_only=True) if r[2] is not None)
    assert ventas_total == float(data["period"]["revenue"])
    assert summary["reports"] == ["Ventas diarias (ingresos cobrados)", "Pedidos", "Existencias por sucursal"]


def test_csv_single_devuelve_csv_y_multiple_zip():
    db, world = make_world()
    payload, media_type, filename, _ = MultiExportService(db).export(["ventas"], "csv")
    assert media_type.startswith("text/csv")
    assert filename.endswith(".csv")
    assert b"# Ventas diarias" in payload

    payload, media_type, filename, _ = MultiExportService(db).export(["ventas", "pedidos"], "csv")
    assert media_type == "application/zip"
    assert filename.endswith(".zip")
    with zipfile.ZipFile(io.BytesIO(payload.getvalue())) as archive:
        assert "criterios.txt" in archive.namelist()
        assert "01_ventas.csv" in archive.namelist()


def test_pdf_se_genera_con_cabecera_y_paginas():
    db, world = make_world()
    payload, media_type, filename, _ = MultiExportService(db).export(["ventas", "pedidos", "pagos"], "pdf")
    assert media_type == "application/pdf"
    assert filename.endswith(".pdf")
    assert payload.getvalue().startswith(b"%PDF")


def test_truncamiento_se_informa_no_se_oculta(monkeypatch):
    from src.infrastructure.config import settings as settings_module

    monkeypatch.setattr(settings_module.settings, "report_export_max_rows", 1)
    db, world = make_world()
    u, b = world["actor"], world["branch"]
    va, pa = world["va"], world["product_a"]
    now = datetime.now(UTC)
    add_order(db, b, u, [line(va, pa, "Camisa, lino", 50, 1)], now - timedelta(days=1), paid_at=now - timedelta(days=1))
    add_order(db, b, u, [line(va, pa, "Camisa, lino", 50, 1)], now - timedelta(days=2), paid_at=now - timedelta(days=2))

    payload, _, _, summary = MultiExportService(db).export(["ventas"], "xlsx")
    assert summary["truncated"] is True
    workbook = load_workbook(io.BytesIO(payload.getvalue()))
    assert "limite maximo de filas" in workbook["Criterios"]["B7"].value


def test_reporte_invalido_es_rechazado_en_el_servicio():
    db, world = make_world()
    with pytest.raises(KeyError):
        MultiExportService(db).export(["reporte_inventado"], "xlsx")


def test_export_soporta_created_at_mixtos_naive_y_aware():
    """Datos historicos pueden quedar naives en SQLite y pedidos nuevos aware;
    el sorteo de export no debe romperse ni omitir filas."""
    db, world = make_world()
    u, b = world["actor"], world["branch"]
    va, pa = world["va"], world["product_a"]
    now = datetime.now(UTC)
    add_order(db, b, u, [line(va, pa, "Camisa, lino", 50, 1)], now - timedelta(days=2), paid_at=now - timedelta(days=2))
    naive = now.replace(tzinfo=None) - timedelta(days=1)
    legacy = OrderModel(number=uuid.uuid4().hex[:24].upper()[:24], user_id=u.id, customer_email="legacy@mail.test",
                        branch_id=b.id, status="paid", payment_status="paid", payment_method="manual",
                        total=Decimal("30"), currency="bob", address={},
                        items=[{"variant_id": str(va.id), "product_id": str(pa.id), "name": "Camisa, lino",
                                "quantity": 1, "unit_price": "30"}],
                        tracking=[], created_at=naive, updated_at=naive)
    db.add(legacy)
    db.commit()
    title, headers, rows = ReportsService(db).export_report("pedidos")
    assert len(rows) == 2
    assert {r["correo"] for r in rows} == {"c@mail.test", "legacy@mail.test"}
    assert title == "Pedidos"